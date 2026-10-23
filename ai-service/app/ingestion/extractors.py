"""Turn an uploaded file into a list of raw text sections (PDF page / slide / document body)."""
from __future__ import annotations

import io
import os
import zipfile

MAX_UNZIPPED_BYTES = 200 * 1024 * 1024  # zip-bomb guard for DOCX / PPTX

SUPPORTED = {".pdf", ".docx", ".pptx", ".txt"}
UNSUPPORTED_HINTS = {
    ".doc": "Old .doc files are not supported. Save it as .docx and upload again.",
    ".ppt": "Old .ppt files are not supported. Save it as .pptx and upload again.",
}


class ExtractionError(Exception):
    """A problem with the uploaded file that is safe to show to the user."""


def file_extension(filename: str) -> str:
    return os.path.splitext(filename.lower())[1]


def validate_upload(filename: str, data: bytes) -> None:
    """Cheap synchronous checks done before accepting an upload."""
    ext = file_extension(filename)
    if ext in UNSUPPORTED_HINTS:
        raise ExtractionError(UNSUPPORTED_HINTS[ext])
    if ext not in SUPPORTED:
        raise ExtractionError("Unsupported file type. Upload a PDF, DOCX, PPTX or TXT file.")
    if not data:
        raise ExtractionError("The uploaded file is empty.")
    if ext == ".pdf" and b"%PDF" not in data[:1024]:
        raise ExtractionError("This file is not a valid PDF.")
    if ext in (".docx", ".pptx"):
        _check_zip(data, ext)


def _check_zip(data: bytes, ext: str) -> None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            total = sum(i.file_size for i in z.infolist())
    except zipfile.BadZipFile:
        raise ExtractionError(f"This file is not a valid {ext.upper().lstrip('.')} document.")
    if total > MAX_UNZIPPED_BYTES:
        raise ExtractionError("This document is too large once unpacked.")


def extract(filename: str, data: bytes) -> list[str]:
    validate_upload(filename, data)
    ext = file_extension(filename)
    try:
        if ext == ".pdf":
            sections = _pdf(data)
        elif ext == ".docx":
            sections = _docx(data)
        elif ext == ".pptx":
            sections = _pptx(data)
        else:
            sections = [_txt(data)]
    except ExtractionError:
        raise
    except Exception as exc:  # parser-specific errors on corrupt files
        raise ExtractionError(f"Could not read this {ext.lstrip('.').upper()} file ({type(exc).__name__}).") from exc

    sections = [s for s in sections if s and s.strip()]
    if not sections:
        raise ExtractionError(
            "No readable text found. Scanned or image-only files are not supported yet (no OCR)."
        )
    return sections


def _pdf(data: bytes) -> list[str]:
    from pypdf import PdfReader

    reader = PdfReader(io.BytesIO(data))
    if reader.is_encrypted:
        try:
            ok = reader.decrypt("")
        except Exception:
            ok = 0
        if not ok:
            raise ExtractionError("This PDF is password-protected.")
    return [(page.extract_text() or "") for page in reader.pages]


def _docx(data: bytes) -> list[str]:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    document = docx.Document(io.BytesIO(data))
    lines: list[str] = []
    for child in document.element.body.iterchildren():
        if child.tag.endswith("}p"):
            lines.append(Paragraph(child, document).text)
        elif child.tag.endswith("}tbl"):
            for row in Table(child, document).rows:
                cells = [c.text.strip() for c in row.cells]
                lines.append(" | ".join(c for c in cells if c))
    return ["\n".join(lines)]


def _pptx(data: bytes) -> list[str]:
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    def shape_texts(shape):
        if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
            for inner in shape.shapes:
                yield from shape_texts(inner)
        elif getattr(shape, "has_table", False) and shape.has_table:
            for row in shape.table.rows:
                cells = [c.text.strip() for c in row.cells]
                yield " | ".join(c for c in cells if c)
        elif shape.has_text_frame:
            yield shape.text_frame.text

    slides: list[str] = []
    for slide in Presentation(io.BytesIO(data)).slides:
        parts = [t for shape in slide.shapes for t in shape_texts(shape)]
        if slide.has_notes_slide:
            notes = slide.notes_slide.notes_text_frame.text
            if notes.strip():
                parts.append(notes)
        slides.append("\n".join(parts))
    return slides


def _txt(data: bytes) -> str:
    if data.startswith((b"\xff\xfe", b"\xfe\xff")):          # UTF-16 only when a BOM says so
        return data.decode("utf-16", errors="replace")
    try:
        return data.decode("utf-8-sig")
    except UnicodeDecodeError:
        return data.decode("cp1252", errors="replace")
